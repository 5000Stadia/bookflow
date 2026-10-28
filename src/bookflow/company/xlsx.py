"""Reading an ``.xlsx`` workbook for imports (R107, Kabe 2026-10-27): the first sheet or a named one, values only.

Through openpyxl in read-only mode with ``data_only`` (each cell's saved value, never its formula), ``keep_vba`` off
(macros are not even copied) and ``keep_links`` off (external links are not read); openpyxl parses its XML through
defusedxml, so entity expansion and external entities are refused. Nothing in the file is executed or evaluated. The
zip is measured before parsing, so a small file that unpacks to gigabytes is refused. A formula whose result was never
saved (a file a program wrote without calculating it) refuses the file, naming its cells, rather than reading blank.

Cells become the text the CSV reader would have seen: numbers in plain decimal notation (``599.56``, never ``5.9956E2``),
whole numbers without ``.0``, dates as ISO (``2026-01-31``), true and false as ``true``/``false``. A number typed into a
cell loses leading zeros in the spreadsheet itself; such columns should be stored as text there.
"""
from __future__ import annotations

import io
import zipfile
from datetime import date, datetime, time
from decimal import Decimal

from bookflow import BookflowError

MAX_UNPACKED = 256 * 1024 * 1024  # bytes the workbook may unpack to
MAX_RATIO = 200  # unpacked over packed, per member above 1 MB
FORMULA_CELLS_SHOWN = 10


def _refuse(message: str, **details) -> BookflowError:
    return BookflowError('E_VALIDATION', message=message, details={'field': 'attachment', **details})


def _check_zip(raw: bytes) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile:
        raise _refuse('The file is a damaged zip, not an .xlsx workbook; save it again as .xlsx or CSV.',
                      problem='not_xlsx') from None
    with archive:
        members = archive.infolist()
        names = {m.filename for m in members}
        if 'xl/workbook.xml' not in names:
            raise _refuse('The file is a zip archive but not an .xlsx workbook (.ods and .xlsb are not read); save it '
                          'as .xlsx or CSV.', problem='not_xlsx')
        if any(m.flag_bits & 0x1 for m in members):
            raise _refuse('The workbook is encrypted; save an unprotected copy as .xlsx or CSV.', problem='encrypted')
        unpacked = sum(m.file_size for m in members)
        if unpacked > MAX_UNPACKED or any(m.file_size > 1024 * 1024 and m.file_size > MAX_RATIO * max(m.compress_size, 1)
                                          for m in members):
            raise _refuse(f'The workbook unpacks to {unpacked:,} bytes, beyond what an import reads; split it or save '
                          'it as CSV.', problem='too_large')


def cell_text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value or value in (float('inf'), float('-inf')):
            return str(value)
        text = format(Decimal(repr(value)), 'f')
        return text[:-2] if text.endswith('.0') else text
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=' ')
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def read(raw: bytes, *, sheet: str | None = None, header_row: bool = True, skip_rows: int = 0):
    from bookflow.company import tabular
    _check_zip(raw)
    import openpyxl
    import openpyxl.xml
    if not openpyxl.xml.DEFUSEDXML:  # never parse an untrusted workbook without the hardened XML parser
        raise BookflowError('E_IO', message='Reading .xlsx needs defusedxml installed beside openpyxl.',
                            details={'check': 'defusedxml'})
    try:
        values = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_vba=False,
                                        keep_links=False)
        formulas = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_vba=False,
                                          keep_links=False)
    except BookflowError:
        raise
    except Exception as exc:  # a malformed workbook: openpyxl raises many kinds
        raise _refuse(f'The workbook cannot be read: {type(exc).__name__}; save it again as .xlsx or CSV.',
                      problem='not_xlsx') from None
    try:
        names = values.sheetnames
        if sheet is None:
            name = names[0]
        elif sheet in names:
            name = sheet
        else:
            raise BookflowError('E_VALIDATION', message=f'The workbook has no sheet {sheet!r}; it has: '
                                + ', '.join(names) + '.', details={'field': 'sheet', 'sheets': names})
        records: list[tuple[int, list[str]]] = []
        unsaved: list[str] = []
        pairs = zip(values[name].iter_rows(values_only=True), formulas[name].iter_rows(values_only=True))
        for line, (row, raw_row) in enumerate(pairs, start=1):
            cells = [cell_text(v) for v in row]
            for i, (v, f) in enumerate(zip(row, raw_row)):
                if v is None and isinstance(f, str) and f.startswith('='):
                    unsaved.append(f'{tabular.column_letter(i)}{line}')
            if any(c.strip() for c in cells):
                records.append((line, cells))
    finally:
        values.close()
        formulas.close()
    if unsaved:
        raise _refuse(f'{len(unsaved)} cell(s) hold a formula whose result was never saved ('
                      + ', '.join(unsaved[:FORMULA_CELLS_SHOWN]) + '); open the file in a spreadsheet program and save '
                      'it, or export it as CSV.', problem='unsaved_formula', cells=unsaved[:FORMULA_CELLS_SHOWN])
    records = records[skip_rows:]
    if not records:
        raise _refuse('The sheet has no rows.', sheet=name)
    if header_row:
        headers = [tabular.clean_text(c)[0].strip() for c in records[0][1]]
        while headers and not headers[-1]:
            headers.pop()
        body = records[1:]
        if not any(headers):
            raise _refuse('The header row is empty.', sheet=name)
        seen = set()
        for header in headers:
            key = tabular.normalize_header(header)
            if key and key in seen:
                raise BookflowError('E_VALIDATION', message=f'The header {header!r} appears twice.',
                                    details={'field': 'content', 'header': header})
            seen.add(key)
    else:
        width = max(len(cells) for _, cells in records)
        headers = [f'Column {tabular.column_letter(i)}' for i in range(width)]
        body = records
    width = len(headers)
    rows, widths = [], {}
    for position, (line, cells) in enumerate(body, start=1):
        extra = cells[width:]
        cells = (cells[:width] + [''] * (width - len(cells)))
        if any(c.strip() for c in extra):  # a value beyond the header's last column
            widths[position] = width + len(extra)
            cells = cells + extra
        rows.append((line, cells))
    return tabular.Table(headers, rows, 'xlsx', False, 'xlsx', tabular.digest(raw), widths, sheet=name)
