"""Kabe's import rules (2026-10-27), retrofitted onto list imports, and reading .xlsx (design/bulk-import.md §2, §3).

- `import check` reads the file once and reports every problem at once, grouped by column, with row numbers and the
  exact cell text, running no command; `preview` and `run` start with it.
- All or nothing by default; `partial` is the opt-in.
- Text keeps its words: a character no text field holds is dropped and reported by row, column and position; what a
  field can hold (accents, scripts, symbols, emoji) is kept. A NUL in a cell no longer refuses the file; a file that is
  not text still does.
- Junk in money, numbers, dates, choices and references refuses the run, never repaired.
- Empty is allowed unless truly needed: the import's required fields are exactly the single command's.
"""
import io

import pytest

from bookflow import BookflowError
from bookflow.company import import_kinds, tabular
from bookflow.core import registry
from tests.test_imports import active_customers, books, company_db, run_all  # noqa: F401


def refusal(fn):
    with pytest.raises(BookflowError) as caught:
        fn()
    return caught.value


def counts_of(books):
    with company_db(books) as db:
        return {t: db.execute(f'SELECT count(*) FROM {t}').fetchone()[0]
                for t in ('customers', 'import_runs', 'import_rows', 'batches', 'audit_events')}


def test_the_check_lists_every_problem_at_once_and_a_refused_preview_changes_nothing(books):
    text = ('Customer,Credit Limit,Terms,Job Start,Phone\n'
            'Ann,10.00,Net 30,,1\n'
            'Ben,$599.♎56,Net 30,2026-01-06,2\n'
            'Cy,12x,Net 99,2026-13-45,3\n'
            'Di,$599.♎56,Net 30,,' + '9' * 80 + '\n'
            ',5,Net 30,,\n'
            'Ann,1,Net 30,,\n')
    before = counts_of(books)
    out = books.run('import check', {'kind': 'customer', 'content': text})
    assert out['ok'] is False and out['rows'] == 6 and out['rows_with_problems'] == 5
    groups = {(g['column'], g['code'], g['value']): g['rows'] for g in out['problems']}
    assert groups == {('Customer', 'no_key', None): [5], ('Customer', 'duplicate_in_file', None): [6],
                      ('Credit Limit', 'bad_value', '$599.♎56'): [2, 4], ('Credit Limit', 'bad_value', '12x'): [3],
                      ('Job Start', 'bad_value', '2026-13-45'): [3], ('Phone', 'E_VALIDATION', '9' * 80): [4]}
    assert "Credit Limit, rows 2, 4: not an amount: '$599.♎56'" in out['summary']
    # Columns in file order; references are looked up and the unknown one asks, once per value.
    assert [g['column'] for g in out['problems']][:3] == ['Customer', 'Customer', 'Credit Limit']
    assert [(u['value'], u['rows']) for u in out['unknown_references']] == [('Net 99', [3])]  # on a refused row too
    assert counts_of(books) == before  # a check changes nothing
    error = refusal(lambda: books.run('import preview', {'kind': 'customer', 'content': text}))
    assert error.code == 'E_VALIDATION' and error.details['check']['problems_total'] == 7
    assert "not an amount: '$599.♎56'" in error.message
    assert counts_of(books) == before  # a refused preview stores no run
    partial = books.run('import preview', {'kind': 'customer', 'content': text, 'partial': True})
    assert partial['counts'] == {'create': 1, 'refused': 5}  # the Net 99 row is refused for its date first


def test_an_incomplete_mapping_is_refused_even_with_partial(books):
    out = books.run('import check', {'kind': 'term', 'content': 'Terms,Net due in days\nNet 45,45\n'})
    assert [(g['code'], g['rows']) for g in out['problems']] == [('unmapped_required', [])]
    assert 'kind' in out['problems'][0]['message']
    error = refusal(lambda: books.run('import preview', {'kind': 'term', 'content': 'Terms\nNet 45\n',
                                                         'partial': True}))
    assert error.code == 'E_VALIDATION' and 'No column maps to kind' in error.message


def test_all_or_nothing_is_the_default_and_partial_the_opt_in(books):
    text = 'Name,Credit Limit\nAlpha,10\nBravo,20\n'
    books.run('customer create', {'name': 'Bravo', 'phone': '1'})
    error = refusal(lambda: books.run('import preview', {'kind': 'customer', 'content': text, 'mode': 'upsert',
                                                         'constants': {'phone': '9' * 70}}))
    assert 'The constant phone: 70 characters' in error.message  # a constant is checked as a cell is
    preview = books.run('import preview', {'kind': 'customer', 'content': text})
    assert preview['run']['atomic'] is True and preview['counts'] == {'create': 1, 'refused': 1}  # Bravo differs
    error = refusal(lambda: books.run('import run', {'run': preview['run']['id']}))
    assert error.code == 'E_VALIDATION' and 'refused no rows' in error.message
    assert active_customers(books) == {'Bravo'}
    ok = books.run('import preview', {'kind': 'customer', 'content': text, 'source': 'other', 'allow': ['differs']})
    assert ok['run']['atomic'] is True
    assert run_all(books, ok['run']['id'])['counts'] == {'created': 1, 'differs': 1, 'not_run': 0}


def test_text_drops_only_what_it_cannot_hold_and_says_where(books):
    text = ('Name,Notes,Company\n'
            '"Johnson, Beck \x00♎️","line one\x07\nline two\tend",Café 東京\x1b\n'
            'Plain,,\n')
    out = books.run('import check', {'kind': 'customer', 'content': text})
    assert out['ok'] is True and out['dropped_total'] == 3
    assert [(d['row'], d['column'], d['position'], d['name']) for d in out['dropped']] == [
        (1, 'Name', 15, 'NUL'), (1, 'Notes', 9, 'BEL'), (1, 'Company', 8, 'ESCAPE')]
    preview = books.run('import preview', {'kind': 'customer', 'content': text})
    assert preview['dropped_total'] == 3 and 'dropped' in ' '.join(preview['notes'])
    run_all(books, preview['run']['id'])
    johnson = books.run('customer show', {'customer': 'Johnson, Beck ♎️'})
    assert johnson['notes'] == 'line one\nline two\tend'  # long text keeps tabs and line breaks
    assert johnson['company_name'] == 'Café 東京'  # accents and other scripts kept
    shown = books.run('import show', {'run': preview['run']['id']})
    assert shown['dropped_total'] == 3 and shown['dropped'][0]['character'] == 'U+0000'
    # A tab is not held by a name: dropped and reported; a NUL in a reference is junk.
    tab = books.run('import check', {'kind': 'customer', 'content': 'Name,Terms\n"Tab\tCo",Net\x0030\n'})
    assert [(d['column'], d['name']) for d in tab['dropped']] == [('Name', 'TAB')]
    assert [(g['column'], g['code']) for g in tab['problems']] == [('Terms', 'bad_value')]


def test_a_file_that_is_not_text_is_refused_and_a_nul_in_a_cell_is_not():
    table = tabular.read(b'Name\nA\x00B\n')
    assert table.rows[0][1] == ['A\x00B']
    for raw in (bytes(range(256)) * 20, b'\x00\x01\x02\x03' * 400, b'%PDF-1.7 ...', b'\xd0\xcf\x11\xe0rest'):
        with pytest.raises(BookflowError) as caught:
            tabular.read(raw)
        assert caught.value.details['problem'] in ('not_text', 'encoding')


def test_blank_cells_and_absent_columns_leave_fields_empty(books):
    text = 'Name,Email,Phone,Credit Limit,Terms\nOnly Name,,,,\nSome,s@x.test,,,Net 30\n'
    assert books.run('import check', {'kind': 'customer', 'content': text})['ok'] is True
    run_all(books, books.run('import preview', {'kind': 'customer', 'content': text})['run']['id'])
    typed = books.run('customer create', {'name': 'Typed'})
    imported = books.run('customer show', {'customer': 'Only Name'})
    volatile = {'id', 'name', 'full_name', 'created_at', 'updated_at', 'created_by', 'updated_by', 'created_via',
                'updated_via', 'version', 'path'}
    same = lambda shown: {k: v for k, v in shown.items() if k not in volatile and not k.endswith('_source_id')}
    assert same(imported) == same(books.run('customer show', {'customer': typed['id']}))
    # An item imports with its name, type and income account (a service item's own rule), the price left blank.
    items = books.run('import preview', {'kind': 'item', 'content': 'Item,Type,Sales Price,Income Account\n'
                                                                  'Bare,service,,Service Income\n'})
    assert items['counts'] == {'create': 1}, items['refused']


@pytest.mark.parametrize('kind', sorted(import_kinds.KINDS))
def test_the_import_requires_exactly_what_the_command_requires(kind):
    model = registry.get(import_kinds.kind(kind).create).input_model
    command = {name for name, info in model.model_fields.items() if info.is_required()} - import_kinds.SUPPLIED
    assert set(import_kinds.required(kind)) == command
    listed = {k['kind']: k for k in __import__('bookflow.company.imports', fromlist=['x']).kinds_list()['kinds']}
    assert set(listed[kind]['required']) == command


def test_a_reference_gone_since_the_preview_refuses_the_all_or_nothing_run(books):
    books.run('term create', {'name': 'Net 21', 'kind': 'standard', 'due_days': 21})
    preview = books.run('import preview', {'kind': 'customer', 'content': 'Name,Terms\nA,Net 21\nB,Net 30\n'})
    books.run('term update', {'term': 'Net 21', 'name': 'Net 22'})
    before = counts_of(books)
    error = refusal(lambda: books.run('import run', {'run': preview['run']['id']}))
    assert error.code == 'E_VALIDATION' and "'Net 21'" in error.message and error.details['check'][0]['rows'] == [1]
    assert counts_of(books) == before
    assert books.run('import show', {'run': preview['run']['id']})['run']['state'] == 'previewed'


# ------------------------------------------------------------------ .xlsx

def _workbook(rows, *, sheet='Customers', extra=None, formula_without_value=False):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    if extra:
        other = wb.create_sheet(extra)
        other.append(['Name'])
        other.append(['From the second sheet'])
    if formula_without_value:
        ws['B3'] = '=1+1'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_an_xlsx_sheet_reads_values_never_formulas(tmp_path):
    from datetime import date
    raw = _workbook([['Customer', 'Credit Limit', 'Since', 'Phone', 'Active'],
                     ['Acme', 599.56, date(2026, 1, 31), '00123', True], ['Bolt', 1250, None, None, False]],
                    extra='Other')
    table = tabular.read(raw)
    assert table.headers == ['Customer', 'Credit Limit', 'Since', 'Phone', 'Active'] and table.sheet == 'Customers'
    assert [cells for _, cells in table.rows] == [['Acme', '599.56', '2026-01-31', '00123', 'true'],
                                                   ['Bolt', '1250', '', '', 'false']]
    assert [line for line, _ in table.rows] == [2, 3] and table.encoding == 'xlsx'
    assert tabular.read(raw, sheet='Other').rows[0][1] == ['From the second sheet']
    with pytest.raises(BookflowError) as caught:
        tabular.read(raw, sheet='Missing')
    assert caught.value.details['sheets'] == ['Customers', 'Other']
    with pytest.raises(BookflowError) as caught:
        tabular.read(_workbook([['Name', 'Amount'], ['A', 1]], formula_without_value=True))
    assert caught.value.details['problem'] == 'unsaved_formula' and caught.value.details['cells'] == ['B3']
    with pytest.raises(BookflowError) as caught:
        tabular.read(b'PK\x03\x04not really a zip')
    assert caught.value.details['problem'] == 'not_xlsx'


def test_an_xlsx_attachment_imports(books, tmp_path):
    import openpyxl.xml
    assert openpyxl.xml.DEFUSEDXML  # the hardened XML parser is what reads every workbook
    company = books.run('company show')
    added = books.run('attachment add', dict(
        record_type='company_info', record_id=company['company_id'], original_filename='customers.xlsx',
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
        input_stream=io.BytesIO(_workbook([['Customer', 'Credit Limit'], ['Sheet Co', 10.5], ['Second Co', 20]])))
    preview = books.run('import preview', {'kind': 'customer', 'attachment': added['attachment']['id']})
    assert preview['counts'] == {'create': 2} and preview['run']['encoding'] == 'xlsx'
    run_all(books, preview['run']['id'])
    assert books.run('customer show', {'customer': 'Sheet Co'})['credit_limit']['amount'] == '10.50'


# ------------------------------------------------------------------ the review of increment 2 (r107-inc2-review.md)

def test_amounts_are_exact_or_refused_never_regrouped(books):
    """High 1: with `.` the decimal mark, `12,50` and `1,5` once landed as 1250.00 and 15.00."""
    out = books.run('import check', {'kind': 'customer', 'decimal': '.',
                                     'content': 'Name,Credit Limit\nA,"1,234.50"\nB,"12,50"\nC,"1,5"\nD,1$5\n'})
    assert {(g['value'], g['code']) for g in out['problems']} == {('12,50', 'bad_value'), ('1,5', 'bad_value'),
                                                                   ('1$5', 'bad_value')}
    assert [g['rows'] for g in out['problems']] == [[2], [3], [4]]


def test_identifiers_refuse_a_stray_character_and_single_line_breaks_become_spaces(books):
    """High 2: an account number `10<NUL>00` became `1000` and updated account 1000; an address line lost the space
    its line break stood for."""
    number = books.run('account query', {'limit': 1})['items'][0]['number']
    junk = number[:2] + '\x00' + number[2:]
    out = books.run('import check', {'kind': 'account', 'mode': 'update',
                                     'content': f'Number,Name\n{junk},Renamed By Import\n'})
    assert [(g['column'], g['code']) for g in out['problems']] == [('Number', 'bad_value')] and not out['dropped']
    email = books.run('import check', {'kind': 'customer', 'content': 'Name,Email\nA,"a\tb@x.test"\n'})
    assert [(g['column'], g['code']) for g in email['problems']] == [('Email', 'bad_value')]
    text = 'Name,Address,Notes\n"Tab\tCo","123 Main St\r\nApt 4","one\ntwo"\n'
    out = books.run('import check', {'kind': 'customer', 'content': text, 'columns': {'Address': 'billing_address.line1'}})
    assert out['ok'] is True
    assert [(d['column'], d['name'], d.get('replaced_with')) for d in out['dropped']] == [
        ('Name', 'TAB', 'SPACE'), ('Address', 'CARRIAGE RETURN', 'SPACE'), ('Address', 'LINE FEED', None)]
    preview = books.run('import preview', {'kind': 'customer', 'content': text,
                                           'columns': {'Address': 'billing_address.line1'}})
    run_all(books, preview['run']['id'])
    shown = books.run('customer show', {'customer': 'Tab Co'})
    assert shown['billing_address']['line1'] == '123 Main St Apt 4' and shown['notes'] == 'one\ntwo'


CHECKED_LIKE_THE_RUN = [
    # (kind, file, mode, extra, the column the check names) -- each once passed the check and was refused by the run.
    ('customer', 'Name,Terms\nAcme,Net 77\n', 'create', {}, 'Terms'),  # an inactive term
    ('customer', 'Name,Parent\nJob A,Old Co\n', 'create', {}, 'Parent'),  # an inactive parent
    ('customer', 'Name,Phone\nSmith: Kitchen,5\n', 'create', {}, 'Name'),  # a colon in a list name
    ('customer', 'Name,Credit Limit\nA,99999999999999999999\n', 'create', {}, 'Credit Limit'),
    ('term', 'Terms,Kind,Net due in days,Discount %\nT1,standard,10,150\n', 'create', {}, 'Discount %'),
    ('term', 'Terms,Kind\nT2,standard\n', 'create', {}, None),  # standard terms need due_days
    ('price-level', 'Price Level,Type,%\nP1,fixed_percent,-200\n', 'create', {}, '%'),
    ('price-level', 'Price Level,Type\nP2,fixed_percent\n', 'create', {}, None),
    ('item', 'Name,Type\nWidget,service\n', 'create', {}, None),  # a service needs its income account
    ('item', 'Name,Type,Income Account\nPart,inventory_part,Service Income\n', 'create', {}, None),
    ('term', 'Terms,Net due in days\nBrand New,10\n', 'upsert', {}, None),  # a new row in upsert needs `kind`
    ('customer', 'Name,Parent\nSelf1,Self1\n', 'create', {}, 'Parent'),  # its own parent
    ('customer', 'Name,Job start\nLoose Job,2026-03-01\n', 'create', {}, 'Job start'),  # a job field, no parent
]


@pytest.mark.parametrize('kind, text, mode, extra, column', CHECKED_LIKE_THE_RUN)
def test_the_check_refuses_what_the_run_would(books, kind, text, mode, extra, column):
    """High 3: the check runs each row through the command's own input model (references as a placeholder id), the
    rules it applies to its input alone, and the books' active flags."""
    books.run('term create', {'name': 'Net 77', 'kind': 'standard', 'due_days': 77})
    books.run('term deactivate', {'term': 'Net 77'})
    books.run('customer create', {'name': 'Old Co'})
    books.run('customer deactivate', {'customer': 'Old Co'})
    out = books.run('import check', {'kind': kind, 'content': text, 'mode': mode, **extra})
    assert out['ok'] is False and out['clean'] is False and out['rows_with_problems'] == 1, out['summary']
    assert {g['column'] for g in out['problems']} == {column}, out['summary']
    error = refusal(lambda: books.run('import preview', {'kind': kind, 'content': text, 'mode': mode, **extra}))
    assert error.code == 'E_VALIDATION'


def test_a_parent_later_in_the_file_runs_first_and_pending_decisions_are_not_ok(books):
    text = 'Name,Parent\nJob 1,Cust 1\nCust 1,\nJob 2,Cust 1\n'
    out = books.run('import check', {'kind': 'customer', 'content': text})
    assert out['ok'] is True and out['unknown_references'] == []
    preview = books.run('import preview', {'kind': 'customer', 'content': text})
    assert preview['counts'] == {'create': 3} and preview['run']['state'] == 'previewed'
    assert run_all(books, preview['run']['id'])['counts'] == {'created': 3, 'not_run': 0}
    assert {'Cust 1', 'Cust 1:Job 1', 'Cust 1:Job 2'} <= active_customers(books)
    # Unknown references and undecided formats are not problems in a cell, but the check is not ok until decided.
    unknown = books.run('import check', {'kind': 'customer', 'content': 'Name,Terms\nZed,Net 99\n'})
    assert (unknown['ok'], unknown['clean']) == (False, True)
    assert [(d['decision'], d['value']) for d in unknown['to_decide']] == [('reference', 'Net 99')]
    ambiguous = books.run('import check', {'kind': 'customer', 'content': 'Name,Credit Limit\nZed,1.234\n'})
    assert (ambiguous['ok'], ambiguous['clean']) == (False, True)
    assert [d['decision'] for d in ambiguous['to_decide']] == ['decimal']


ROW_PARTS = {
    'Name': ['Ann', 'Bo\x00b', 'Smith: Kitchen', 'Cy\tDe', ''],
    'Credit Limit': ['10', '12,50', '1,234.50', '$5', '1$5', '99999999999999999999', '(3.00)', '', '1.005'],
    'Terms': ['Net 30', 'Net 77', 'Net 99', 'Net\x0030', ''],
    'Email': ['a@x.test', 'a\x00@x.test', ''],
    'Phone': ['555 0100', '9' * 80, ''],
}


def test_property_a_row_the_check_passes_is_never_refused_for_a_cell(books):
    """High 3, as a property: over generated rows mixing the tricky cells above, whenever the check finds no problem
    the preview refuses no row, and the run lands every row."""
    import itertools
    import random
    books.run('term create', {'name': 'Net 77', 'kind': 'standard', 'due_days': 77})
    books.run('term deactivate', {'term': 'Net 77'})
    rng = random.Random(107)
    combos = list(itertools.product(*ROW_PARTS.values()))
    passed = 0
    for n, combo in enumerate(rng.sample(combos, 40)):
        name = combo[0] and f'{combo[0]} {n}'
        cells = [name, *combo[1:]]
        text = ','.join(ROW_PARTS) + '\n' + ','.join('"' + c + '"' for c in cells) + '\n'
        check = books.run('import check', {'kind': 'customer', 'content': text})
        if not check['clean']:
            continue
        preview = books.run('import preview', {'kind': 'customer', 'content': text, 'partial': True})
        assert 'refused' not in preview['counts'], (cells, preview['refused'])
        if check['ok']:
            passed += 1
            assert run_all(books, preview['run']['id'])['counts']['created'] == 1, cells
    assert passed >= 3  # the sample holds rows that pass, not only refusals


def test_all_or_nothing_refuses_differs_and_skipped_unless_allowed(books):
    """Medium 5a: an atomic run once ended `done` with `differs` and `skipped` rows unapplied."""
    books.run('customer create', {'name': 'Acme', 'phone': '1'})
    text = 'Name,Phone\nAcme,2\nNew Co,3\n'
    preview = books.run('import preview', {'kind': 'customer', 'content': text})
    assert preview['counts'] == {'create': 1, 'refused': 1} and preview['refused'][0]['code'] == 'differs'
    assert 'allow' in refusal(lambda: books.run('import run', {'run': preview['run']['id']})).message + \
        preview['refused'][0]['message']
    allowed = books.run('import preview', {'kind': 'customer', 'content': text, 'allow': ['differs']})
    assert allowed['counts'] == {'create': 1, 'differs': 1}
    assert run_all(books, allowed['run']['id'])['counts'] == {'created': 1, 'differs': 1, 'not_run': 0}
    update = books.run('import preview', {'kind': 'customer', 'mode': 'update', 'content': 'Name,Phone\nNobody,1\n'})
    assert update['counts'] == {'refused': 1} and update['refused'][0]['code'] == 'missing'


def test_cancelling_an_atomic_run_undoes_what_it_made(books):
    """Medium 5b: `import cancel` between slices once left the created rows and a decided reference in the books."""
    preview = books.run('import preview', {'kind': 'customer',
                                           'content': 'Name,Customer Type\nAcme,Gold\nBeta,Gold\nCy,Gold\n'})
    books.run('import decide', {'run': preview['run']['id'], 'create': ['Gold']})
    first = books.run('import run', {'run': preview['run']['id'], 'limit': 1})
    assert first['run']['state'] == 'running' and 'Acme' in active_customers(books)
    cancelled = books.run('import cancel', {'run': preview['run']['id']})
    assert cancelled['run']['state'] == 'cancelled' and 'undone' in ' '.join(cancelled['notes'])
    assert not {'Acme', 'Beta', 'Cy'} & active_customers(books)
    with company_db(books) as db:
        assert db.execute("SELECT active FROM customer_types WHERE name = 'Gold'").fetchall() in ([], [(0,)])


def test_constants_and_saved_defaults_are_checked_like_cells(books):
    out = books.run('import check', {'kind': 'customer', 'content': 'Name\nA\n',
                                     'constants': {'credit_limit': '1$5', 'email': 'a\x00@x.test'}})
    assert sorted(g['code'] for g in out['problems']) == ['bad_value', 'bad_value']
    assert all(g['rows'] == [] and g['message'].startswith('The constant') for g in out['problems'])
    good = books.run('import check', {'kind': 'customer', 'content': 'Name\nA\n', 'constants': {'credit_limit': '5'}})
    assert good['ok'] is True and good['mapping']['constants'] == {'credit_limit': '5.00'}


# ------------------------------------------------------------------ .xlsx, after the review (High 4)

def _patched(raw, old, new):
    import zipfile
    source, out = zipfile.ZipFile(io.BytesIO(raw)), io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                text = data.decode()
                assert old in text, old
                data = text.replace(old, new).encode()
            target.writestr(info.filename, data)
    return out.getvalue()


def test_an_xlsx_sheet_is_read_from_its_cells_once_and_bounded(monkeypatch):
    import time
    raw = _workbook([['Name', 'Amt'], ['A', 1], ['B', 2]])
    # A wrong or absent <dimension> neither hides rows nor pads empty ones.
    for dimension in ('<dimension ref="A1" />', ''):
        table = tabular.read(_patched(raw, '<dimension ref="A1:B3" />', dimension))
        assert [cells for _, cells in table.rows] == [['A', '1'], ['B', '2']]
    # The 4.9 KB file that once took 14 s: a declared A1:XFD2000 and one row far below.
    started = time.monotonic()
    far = _patched(raw, '<dimension ref="A1:B3" />', '<dimension ref="A1:XFD2000" />')
    far = _patched(far, '</sheetData>', '<row r="2000"><c r="B2000" t="inlineStr"><is><t>9</t></is></c></row>'
                                        '</sheetData>')
    assert [line for line, _ in tabular.read(far).rows] == [2, 3, 2000] and time.monotonic() - started < 2
    wide = _patched(raw, '</sheetData>', '<row r="9"><c r="XFD9" t="inlineStr"><is><t>Z</t></is></c></row>'
                                         '</sheetData>')
    with pytest.raises(BookflowError) as caught:
        tabular.read(wide)
    assert caught.value.details['problem'] == 'too_wide'
    monkeypatch.setattr(tabular, 'MAX_ROWS', 1)
    with pytest.raises(BookflowError) as caught:
        tabular.read(raw)
    assert caught.value.details['problem'] == 'too_many_rows'
    with pytest.raises(BookflowError) as caught:
        tabular.read('Name\nA\nB\n')
    assert caught.value.details['problem'] == 'too_many_rows'  # the same cap as a text file


def test_an_xlsx_array_formula_without_a_value_an_error_value_and_damaged_xml_refuse():
    raw = _workbook([['Name', 'Amt'], ['A', 1], ['B', 2]])
    array = _patched(raw, '<c r="B3" t="n"><v>2</v></c>', '<c r="B3"><f t="array" ref="B3">SUM(1,2)</f></c>')
    with pytest.raises(BookflowError) as caught:
        tabular.read(array)
    assert caught.value.details['problem'] == 'unsaved_formula' and caught.value.details['cells'] == ['B3']
    cached = _patched(raw, '<c r="B2" t="n"><v>1</v></c>', '<c r="B2"><f>1+1</f><v>2</v></c>')
    assert tabular.read(cached).rows[0][1] == ['A', '2']  # a saved value is read, the formula never run
    error = _patched(raw, '<c r="A2" t="inlineStr"><is><t>A</t></is></c>', '<c r="A2" t="e"><v>#REF!</v></c>')
    with pytest.raises(BookflowError) as caught:
        tabular.read(error)
    assert caught.value.details['problem'] == 'error_value' and '#REF!' in caught.value.message
    damaged = _patched(raw, '</sheetData>', '<row r="4"><c r="A4"></sheetData>')
    with pytest.raises(BookflowError) as caught:
        tabular.read(damaged)
    assert caught.value.details['problem'] == 'not_xlsx' and 'damaged' in caught.value.message
